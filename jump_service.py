"""CRUD e importação transacional das coletas de salto."""

from __future__ import annotations

import io
import math
import re
from contextlib import AbstractContextManager
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping
from unicodedata import combining, normalize

from openpyxl import load_workbook

from athlete_matching import (
    athlete_display_name,
    athlete_selection_label,
    unique_matching_athlete_id,
)
from database import database_connection, database_write_connection

JUMP_GROUP = "Saltos"
JUMP_METRICS = (
    "CMJ1", "CMJ2", "CMJ3", "MAIOR_CMJ",
    "SJ1", "SJ2", "SJ3", "MAIOR_SJ",
)
MAX_WORKBOOK_SIZE = 20 * 1024 * 1024
IMPORT_LOCK_ID = 1_947_703_002
ConnectionFactory = Callable[[], AbstractContextManager[Any]]


class DuplicateJumpCollectionError(ValueError):
    """Indica que o atleta já possui uma coleta de salto na data."""


@dataclass(frozen=True)
class JumpCollection:
    athlete_id: int
    collected_at: date | datetime
    cmj1: float | None = None
    cmj2: float | None = None
    cmj3: float | None = None
    maior_cmj: float | None = None
    sj1: float | None = None
    sj2: float | None = None
    sj3: float | None = None
    maior_sj: float | None = None


@dataclass(frozen=True)
class PreparedJumpRow:
    filename: str
    sheet: str
    row_number: int
    athlete_id: int | None
    athlete: str
    recognized_name: str
    collected_at: date | None
    measurements: Mapping[str, float]
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class JumpImportPreview:
    row: PreparedJumpRow
    status: str
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class JumpImportResult:
    inserted_collections: int
    inserted_measurements: int
    duplicate_collections: int
    conflicts: int


def _normalized_header(value: object) -> str:
    text = "" if value is None else str(value).strip().casefold()
    decomposed = normalize("NFKD", text)
    return "".join(
        character for character in decomposed
        if not combining(character) and (character.isalnum() or character == "_")
    ).upper()


def _positive_measurement(value: object, label: str) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.casefold() in {"s/d", "sd", "d/", "n/a", "none"}:
        return None
    try:
        number = float(text.replace(",", "."))
    except (TypeError, ValueError) as error:
        raise ValueError(f"{label} deve ser numérico ou S/D.") from error
    if not math.isfinite(number):
        raise ValueError(f"{label} deve ser finito.")
    if number <= 0:
        return None
    return number


def _timestamp(value: date | datetime) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    raise ValueError("Data da coleta é obrigatória.")


def collection_measurements(collection: JumpCollection) -> dict[str, float]:
    values = {
        "CMJ1": collection.cmj1, "CMJ2": collection.cmj2,
        "CMJ3": collection.cmj3, "MAIOR_CMJ": collection.maior_cmj,
        "SJ1": collection.sj1, "SJ2": collection.sj2,
        "SJ3": collection.sj3, "MAIOR_SJ": collection.maior_sj,
    }
    measurements = {
        name: number
        for name, value in values.items()
        if (number := _positive_measurement(value, name)) is not None
    }
    if not measurements:
        raise ValueError("Informe ao menos uma medida válida de CMJ ou SJ.")
    return measurements


def _metric_ids(cursor: Any) -> dict[str, int]:
    cursor.execute(
        """
        SELECT m.id_medida, UPPER(TRIM(m.nome))
        FROM public.medida m
        JOIN public.grupo_medida gm
          ON gm.id_grupo_medida = m.id_grupo_medida
        WHERE LOWER(TRIM(gm.nome)) = LOWER(%s)
          AND UPPER(TRIM(m.nome)) = ANY(%s)
        """,
        (JUMP_GROUP, list(JUMP_METRICS)),
    )
    result = {str(name): int(metric_id) for metric_id, name in cursor.fetchall()}
    missing = sorted(set(JUMP_METRICS) - set(result))
    if missing:
        raise RuntimeError("Medidas de salto ausentes no banco: " + ", ".join(missing))
    return result


def _lock_athlete(cursor: Any, athlete_id: int) -> None:
    cursor.execute(
        "SELECT id_atleta FROM public.atleta WHERE id_atleta = %s FOR UPDATE",
        (int(athlete_id),),
    )
    if cursor.fetchone() is None:
        raise ValueError("Jogador não encontrado no banco.")


def _existing_measurements(
    cursor: Any, athlete_id: int, collected_at: date | datetime
) -> dict[str, float]:
    cursor.execute(
        """
        SELECT UPPER(TRIM(m.nome)), mv.valor
        FROM public.medida_valor mv
        JOIN public.medida m ON m.id_medida = mv.id_medida
        JOIN public.grupo_medida gm
          ON gm.id_grupo_medida = m.id_grupo_medida
        WHERE mv.id_atleta = %s AND mv.data::date = %s
          AND LOWER(TRIM(gm.nome)) = LOWER(%s)
          AND UPPER(TRIM(m.nome)) = ANY(%s)
        """,
        (int(athlete_id), _timestamp(collected_at).date(), JUMP_GROUP, list(JUMP_METRICS)),
    )
    return {
        str(name): float(value)
        for name, value in cursor.fetchall()
        if value is not None
    }


def _insert_measurements(
    cursor: Any,
    athlete_id: int,
    collected_at: date | datetime,
    measurements: Mapping[str, float],
    metric_ids: Mapping[str, int],
) -> int:
    timestamp = _timestamp(collected_at)
    values = [
        (int(athlete_id), None, metric_ids[name], value, str(value), timestamp)
        for name, value in measurements.items()
    ]
    cursor.executemany(
        """
        INSERT INTO public.medida_valor
            (id_atleta, id_partida, id_medida, valor, valor_texto, data)
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        values,
    )
    return len(values)


def create_jump_collection(
    collection: JumpCollection,
    *,
    connection_factory: ConnectionFactory = database_write_connection,
) -> int:
    measurements = collection_measurements(collection)
    with connection_factory() as db, db.cursor() as cursor:
        _lock_athlete(cursor, collection.athlete_id)
        metric_ids = _metric_ids(cursor)
        if _existing_measurements(cursor, collection.athlete_id, collection.collected_at):
            raise DuplicateJumpCollectionError(
                "Já existe uma coleta de salto para esse jogador nessa data."
            )
        return _insert_measurements(
            cursor, collection.athlete_id, collection.collected_at,
            measurements, metric_ids,
        )


def update_jump_collection(
    original_athlete_id: int,
    original_collected_at: date | datetime,
    collection: JumpCollection,
    *,
    connection_factory: ConnectionFactory = database_write_connection,
) -> int:
    measurements = collection_measurements(collection)
    with connection_factory() as db, db.cursor() as cursor:
        _lock_athlete(cursor, original_athlete_id)
        if collection.athlete_id != int(original_athlete_id):
            _lock_athlete(cursor, collection.athlete_id)
        metric_ids = _metric_ids(cursor)
        existing = _existing_measurements(
            cursor, original_athlete_id, original_collected_at
        )
        if not existing:
            raise ValueError("Coleta de salto não encontrada.")
        target_changed = (
            int(original_athlete_id) != int(collection.athlete_id)
            or _timestamp(original_collected_at).date()
            != _timestamp(collection.collected_at).date()
        )
        if target_changed and _existing_measurements(
            cursor, collection.athlete_id, collection.collected_at
        ):
            raise DuplicateJumpCollectionError(
                "Já existe uma coleta de salto para o jogador e data selecionados."
            )
        cursor.execute(
            """
            DELETE FROM public.medida_valor mv
            USING public.medida m, public.grupo_medida gm
            WHERE mv.id_medida = m.id_medida
              AND m.id_grupo_medida = gm.id_grupo_medida
              AND mv.id_atleta = %s AND mv.data::date = %s
              AND LOWER(TRIM(gm.nome)) = LOWER(%s)
              AND UPPER(TRIM(m.nome)) = ANY(%s)
            """,
            (int(original_athlete_id), _timestamp(original_collected_at).date(),
             JUMP_GROUP, list(JUMP_METRICS)),
        )
        return _insert_measurements(
            cursor, collection.athlete_id, collection.collected_at,
            measurements, metric_ids,
        )


def delete_jump_collection(
    athlete_id: int,
    collected_at: date | datetime,
    *,
    connection_factory: ConnectionFactory = database_write_connection,
) -> int:
    with connection_factory() as db, db.cursor() as cursor:
        _lock_athlete(cursor, athlete_id)
        cursor.execute(
            """
            DELETE FROM public.medida_valor mv
            USING public.medida m, public.grupo_medida gm
            WHERE mv.id_medida = m.id_medida
              AND m.id_grupo_medida = gm.id_grupo_medida
              AND mv.id_atleta = %s AND mv.data::date = %s
              AND LOWER(TRIM(gm.nome)) = LOWER(%s)
              AND UPPER(TRIM(m.nome)) = ANY(%s)
            RETURNING mv.id_medida_valor
            """,
            (int(athlete_id), _timestamp(collected_at).date(), JUMP_GROUP,
             list(JUMP_METRICS)),
        )
        deleted = len(cursor.fetchall())
        if deleted == 0:
            raise ValueError("Coleta de salto não encontrada.")
        return deleted


def _header_columns(values: Iterable[object]) -> dict[str, int]:
    columns: dict[str, int] = {}
    repeated_sj = 0
    for index, value in enumerate(values):
        header = _normalized_header(value)
        if index == 0 and header == "A":
            header = "NOME"
        if header == "SJ":
            repeated_sj += 1
            header = f"SJ{repeated_sj}"
        if header in {"NOME", "APELIDO", *JUMP_METRICS}:
            columns[header] = index
    return columns


def _sheet_date(title: str) -> date | None:
    if not re.fullmatch(r"\d{8}", title.strip()):
        return None
    try:
        return datetime.strptime(title.strip(), "%d%m%Y").date()
    except ValueError:
        return None


def extract_jump_workbook(
    content: bytes,
    filename: str,
    athletes: list[dict[str, object]],
) -> dict[str, object]:
    """Extrai as abas datadas do XLSX sem escrever no banco."""
    if not content:
        raise ValueError("A planilha enviada está vazia.")
    if len(content) > MAX_WORKBOOK_SIZE:
        raise ValueError("A planilha deve ter no máximo 20 MB.")
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as error:
        raise ValueError("O arquivo enviado não é um XLSX válido.") from error

    rows: list[dict[str, object]] = []
    skipped_sheets: list[str] = []
    errors: list[str] = []
    for sheet in workbook.worksheets:
        collected_at = _sheet_date(sheet.title)
        if collected_at is None:
            skipped_sheets.append(sheet.title)
            continue
        header_row = None
        columns: dict[str, int] = {}
        for row_number in range(1, min(sheet.max_row, 10) + 1):
            candidate = _header_columns(
                sheet.cell(row_number, column).value
                for column in range(1, sheet.max_column + 1)
            )
            if "NOME" in candidate and "CMJ1" in candidate:
                header_row, columns = row_number, candidate
                break
        if header_row is None:
            errors.append(f"{sheet.title}: cabeçalho de saltos não encontrado.")
            continue

        for row_number in range(header_row + 1, sheet.max_row + 1):
            raw_name = sheet.cell(row_number, columns["NOME"] + 1).value
            raw_nickname = (
                sheet.cell(row_number, columns["APELIDO"] + 1).value
                if "APELIDO" in columns else None
            )
            name = str(raw_name or "").strip()
            nickname = str(raw_nickname or "").strip()
            if not name and not nickname:
                continue
            if _normalized_header(name or nickname) in {"MEDIA", "AVERAGE", "AVERAGES"}:
                continue

            row_errors: list[str] = []
            measurements: dict[str, float] = {}
            for metric in JUMP_METRICS:
                if metric not in columns:
                    continue
                raw_value = sheet.cell(row_number, columns[metric] + 1).value
                try:
                    value = _positive_measurement(raw_value, metric)
                except ValueError as error:
                    row_errors.append(str(error))
                    continue
                if value is not None:
                    measurements[metric] = value
            if not measurements:
                continue

            name_id = unique_matching_athlete_id(name, athletes) if name else None
            nickname_id = (
                unique_matching_athlete_id(nickname, athletes) if nickname else None
            )
            athlete_id = name_id or nickname_id
            if name_id and nickname_id and name_id != nickname_id:
                athlete_id = None
                row_errors.append("Nome e apelido correspondem a jogadores diferentes.")
            recognized = nickname or name
            athlete = next(
                (item for item in athletes if int(item["id_atleta"]) == athlete_id),
                None,
            )
            if athlete_id is None:
                row_errors.append(f'Jogador "{recognized}" não identificado unicamente.')
            rows.append({
                "arquivo": Path(filename).name,
                "aba": sheet.title,
                "linha": row_number,
                "data_coleta": collected_at,
                "nome_reconhecido": recognized,
                "athlete_id": athlete_id,
                "jogador": athlete_selection_label(athlete) if athlete else None,
                **{metric.lower(): measurements.get(metric) for metric in JUMP_METRICS},
                "erros": row_errors,
            })
    workbook.close()
    return {
        "arquivo": Path(filename).name,
        "linhas": rows,
        "abas_ignoradas": skipped_sheets,
        "erros": errors,
    }


def prepare_jump_rows(payload: list[dict[str, object]]) -> list[PreparedJumpRow]:
    prepared: list[PreparedJumpRow] = []
    seen: set[tuple[int, date]] = set()
    for item in payload:
        errors = [str(error) for error in item.get("erros", [])]
        try:
            athlete_id = int(item["athlete_id"]) if item.get("athlete_id") else None
        except (TypeError, ValueError):
            athlete_id = None
        if athlete_id is None:
            errors.append("Selecione um jogador cadastrado.")
        raw_date = item.get("data_coleta")
        if isinstance(raw_date, datetime):
            collected_at = raw_date.date()
        elif isinstance(raw_date, date):
            collected_at = raw_date
        else:
            try:
                collected_at = date.fromisoformat(str(raw_date))
            except (TypeError, ValueError):
                collected_at = None
                errors.append("Data da coleta inválida.")
        measurements: dict[str, float] = {}
        for metric in JUMP_METRICS:
            try:
                value = _positive_measurement(item.get(metric.lower()), metric)
            except ValueError as error:
                errors.append(str(error))
                continue
            if value is not None:
                measurements[metric] = value
        if not measurements:
            errors.append("Informe ao menos uma medida válida.")
        if athlete_id is not None and collected_at is not None:
            key = (athlete_id, collected_at)
            if key in seen:
                errors.append("Jogador e data repetidos no lote.")
            seen.add(key)
        prepared.append(PreparedJumpRow(
            filename=str(item.get("arquivo") or "planilha.xlsx"),
            sheet=str(item.get("aba") or ""),
            row_number=int(item.get("linha") or 0),
            athlete_id=athlete_id,
            athlete=str(item.get("jogador") or ""),
            recognized_name=str(item.get("nome_reconhecido") or ""),
            collected_at=collected_at,
            measurements=measurements,
            errors=tuple(dict.fromkeys(errors)),
        ))
    return prepared


def _same_measurements(first: Mapping[str, float], second: Mapping[str, float]) -> bool:
    return set(first) == set(second) and all(
        math.isclose(first[name], second[name], rel_tol=0, abs_tol=1e-9)
        for name in first
    )


def preview_jump_import(
    rows: list[PreparedJumpRow],
    *,
    connection_factory: ConnectionFactory = database_connection,
) -> list[JumpImportPreview]:
    previews: list[JumpImportPreview] = []
    with connection_factory() as db, db.cursor() as cursor:
        pending: list[PreparedJumpRow] = []
        for row in rows:
            if row.errors or row.athlete_id is None or row.collected_at is None:
                previews.append(JumpImportPreview(row, "erro", row.errors))
                continue
            existing = _existing_measurements(
                cursor, row.athlete_id, row.collected_at
            )
            status = (
                "novo" if not existing
                else "duplicado" if _same_measurements(existing, row.measurements)
                else "conflito"
            )
            errors = (
                ("Já existe uma coleta diferente para esse jogador e data.",)
                if status == "conflito" else ()
            )
            previews.append(JumpImportPreview(row, status, errors))
    return previews


def import_jump_rows(
    rows: list[PreparedJumpRow],
    *,
    connection_factory: ConnectionFactory = database_write_connection,
) -> JumpImportResult:
    if not rows:
        raise ValueError("Não há coletas de salto para importar.")
    if any(row.errors for row in rows):
        raise ValueError("O lote possui linhas inválidas.")
    inserted_collections = inserted_measurements = duplicates = conflicts = 0
    with connection_factory() as db, db.cursor() as cursor:
        cursor.execute("SELECT pg_advisory_xact_lock(%s)", (IMPORT_LOCK_ID,))
        metric_ids = _metric_ids(cursor)
        pending: list[PreparedJumpRow] = []
        for row in rows:
            if row.athlete_id is None or row.collected_at is None:
                raise ValueError("O lote possui jogador ou data inválidos.")
            _lock_athlete(cursor, row.athlete_id)
            existing = _existing_measurements(cursor, row.athlete_id, row.collected_at)
            if existing:
                if _same_measurements(existing, row.measurements):
                    duplicates += 1
                    continue
                raise ValueError(
                    "O lote possui conflito com uma coleta alterada no banco. "
                    "Valide novamente antes de enviar."
                )
            pending.append(row)
        for row in pending:
            inserted_measurements += _insert_measurements(
                cursor, row.athlete_id, row.collected_at,
                row.measurements, metric_ids,
            )
            inserted_collections += 1
    return JumpImportResult(
        inserted_collections, inserted_measurements, duplicates, conflicts
    )
